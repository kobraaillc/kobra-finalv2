import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from app.schemas.pleading import PleadingPaperPayload, PleadingJurisdiction

class PleadingRendererEngine:
    def __init__(self, payload: PleadingPaperPayload):
        self.payload = payload
        self.doc = docx.Document()
        self._configure_styles()

    def _configure_styles(self):
        normal_style = self.doc.styles['Normal']
        normal_style.font.name = 'Times New Roman'
        normal_style.font.size = Pt(12)
        normal_style.font.color.rgb = RGBColor(0,0,0)

    def generate(self) -> docx.Document:
        if self.payload.jurisdiction == PleadingJurisdiction.CA_SUPERIOR:
            return self._build_california_pleading()
        elif self.payload.jurisdiction == PleadingJurisdiction.FL.CIRCUIT:
            return self._build_florida_pleading()
        else:
            raise ValueError(f"Unsupported Jurisdiction: {self.payload.jurisdiction}")

    def _build_california_Pleading(self) -> docx.Document:
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.25)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)

        attorney = self.payload.attorney_of_record
        court = self.payload.court_info
        parties = self.payload.parties
        meta = self.payload.case_metadata

        plaintiff_name = parties.plaintiffss[0].upper() if parties.plaintiffs else "PLAINTIFF"
        attorney_lines = [
            f"{attorney.attorney_name.upper()} (SBN {attorney.bar_number})",
            f"{attorney.firm_name}",
            f"{attorney.address_line1}",
            f"{attorney.city_state_zip}",
            f"Telephone: {attorney.phone}",
            f"Email: {attorney.primary_email}",
            "",
            f"Attorneys for Plaintiff {plaintiff_name}"
        ]
        for line in attorney_lines:
            p = self.doc.add_paragraph(line)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)

        self.doc.add_paragraph("")

        caption_table = self.doc.add_table(rows=2, cols=2)
        caption_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        caption_table.autofit = False

        widths = [Inches(3.25), Inches(3.0)]
        for row in caption_table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w

        cell_00 = caption_table.cell(0,0)
        cell_01 = caption_table.cell(0,1)
        cell_00.merge(cell_01)
        p_court = cell_00.paragraphs[0]
        p_court.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_court.add_run(f"{court.court_name.upper()}\n{court.county_or_circuit.upper()}").bold = True
        p_court.paragraph_format.space_after = Pt(12)

        p_parties = caption_table.cell(1,0).paragraphs[0]
        plaintiff_str = ", ".join(parties, self.payload.parties.plaintiffs)
        defendant_str =", ".join(parties, self.payload.parties.defendants)
        p_parties.add_run(f"{plaintiff_str},\n\nv.\n\n{defendant_str},\n\n\tDefendants.")
        p_parties.paragraph_format.line_spacing = 1.15

        p_meta = caption_table.cell(1,1).paragraphs[0]
        p_meta.add_run(f"CASE NO. {meta.case_number}\n\n").bold = True
        p_meta.add_run(f"{self.payload.document_title.upper()}\n]n").bold = True
        p_meta.add_run(f"Dept: {court.division_dept or 'N/A'}\nJudge: {meta.judge_name or 'N/A'}")

        tcPr = caption_table.cell(1,0)._tc.get_or_add_tcPr()
        borders = parse_xml(r'<w:tcBorders %s><w:right w:val="single" w:sz="12" w:space="0" w:color="000000"/></w:tcBorders>' % nsdecls('w'))
        tcPr.append(borders)

        return self.doc

    def _build_florida_pleading(self) -> docx.Document:
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        court = self.payload.court_info
        parties = self.payload.parties
        meta = self.payload.case_metadata
        attorney = self.payload.attorney_of_record

        p_stamp = self.doc.add_paragraph()
        p_stamp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_stamp = p_stamp.add_run("[RESERVED FOR E-FILING STAMP -3\" x 3\"]")
        r_stamp.font_size = Pt(8)
        r_stamp.font.italice = True
        r_stamp.font.color.rgb = RGBColor(128, 128, 128)

        p_court = self.doc.add_paragraph()
        p_court.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_court.add_run(
            f"{court.court_name.upper()}\n"
            f"{court.county_or_circuit.upper()}\n\n"
            f"CIVIL DIVISION"
        ).bold = True
        p_court.paragraph_format.space_after = Pt(18)

        caption_table = self.doc.add_table(rows=1, cols=2)
        caption_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        caption_table.autofit = False

        widths = [Inches(3.5), Inches(3.0)]
        for i, w in enumerate(widths):
            caption_table.rows[0].cells[i].width = w

        p_parties = caption_table.cell(0,0).paragraphs[0]
        plaintiff_str = ", ".join(parties.plaintiffs)
        defendant_str = ", ".join(parties.defendants)
        p_parties.add_run(f"{plaintiff_str},\n\n\tPlaintiff,\n\nv.n\n\n{defendant_str},\n\n\tDefendant.\n_________________________________/")

        p_meta = caption_table.cell(0.1).paragraphs[0]
        p_meta.add_run(f"CASE NO.: {meta.case_number}\n").bold = True
        p_meta.add_run(f"DIVISION: {court.division_dept or '08'}").bold = True

        p_title =self.doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.add_run(self.payload.document_title.upper()).bold = True
        p_title.paragraph_format.space_before = Pt(18)
        p_title.paragraph_format.space_after = Pt(18)

        p_sig = self.doc.add_paragraph()
        p_sig.paragraph_format.line_spacing = 1.15
        p_sig.paragraph_format.space_before = Pt(24)
        p_sig.add_run("Respectfully submitted, \n\n")
        p_sig.add_run(f"{attorney.firm_name}\n").bold = True
        p_sig.add_run("Attorneys for Plaintiff\n")
        p_sig.add_run(f"{attorney.address_line1}\n{attorney.city_state_zip}\n")
        p_sig.add_run(f"Telephone: {attorney.phone}\n\n")
        p_sig.add_run(f"By: /s/ {attorney.attorney_name}\n").bold = True
        p_sig.add_run(f"Florida Bar No.: {attorney.bar_number}\n")
        p_sig.add_run(f"Primary Email: {attorney.primary_email}\n")
        if attorney.secondary_email:
            p_sig.add_run(f"Secondary Email: {attorney.secondary_email}")

        return self .doc


async def render_pleading_document(payload: PleadingPaperPayload, llm_body_text:str) -> docx.Document:
    """
    Instantiates the renderer engine and append generated LLM body text
    into the formated pleading document
    """

    engine = PleadingRendererEngine(payload=payload)
    doc = engine.generate()


    doc.add_paragraph(llm_body_text)

    return doc