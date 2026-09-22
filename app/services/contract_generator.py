import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.schemas.contract import ContractPayload, ContractTypeEnum

class ContractRendererEngine:
    def __init__(self, payload: ContractPayload):
        self.payload = payload
        self.doc = docx.Document()
        self._configure_document_styles()

    def _configure_document_styles(self):
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        normal_style = self.doc.styles['Normal']
        normal_style.font.name = 'Times New Roman'
        normal_style.font.size = Pt(11)
        normal_style.font.color.rgb = RGBColor(0,0,0)

    def generate(self) -> docx.Document:
        if self.payload.contract_type == ContractTypeEnum.NDA:
            return self ._build_nda_agreement()
        else:
            return self ._build_generic_agreement()

    def _build_nda_agreement(self) -> docx.Document:
        p_title = self.doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.add_run(self.payload.contract_type.value.upper()).bold = True
        p_title.paragraph_format.space_after = Pt(18)

        p_preamble = self.doc.add_paragaph()
        p_preamble.paragraph_format.line_spacing = 1.15
        p_preamble.paragraph_format.space_after = Pt(12)

        disc = self.payload.disclosing_party
        recv = self.payload.receiving_party
        terms = self.payload.terms

        preamble_text = (
            f"This Non-Disclosure Agreement (\"Agreement\") is entered into and made effective as of "
            f"{terms.effective_date.strftime('%B %d, %Y')} (\"Effective Date\"), by and between "
            f"{disc.legal_name}, a {disc.entity_type.value} (\"Disclosing Party\") and "
            f"{recv.legal_name}, a {recv.entity_type.value} (\"Receiving Party\"). "
            f"Disclosing Party and Receiving Party may collectively be referred to as the \"Parties.\""
        )
        p_preamble.add_run(preamble_text)

        p_rec = self.doc.add_paragraph()
        p_rec.add_run("RECITALS").bold = True
        p_rec.paragraph_format.space_before = Pt(12)
        p_rec.paragraph_format.space_after = Pt(6)

        recital_text = (
            f"WHEREAS, the Parties wish to engage in discussions regarding {self.payload.purpose_description} "
            f"(the \"Purpose\"), and in connection with such Purpose, Disclosing Party may disclose confidential "
            f"and proprietary information to Receiving Party."
        )
        p_r_body = self.doc.add_paragraph(recital_text)
        p_r_body.paragraph_format.line_spacing = 1.15
        p_r_body.paragraph_format.space_after = Pt(12)

        clauses = [
            ("1. Confidential Information.", f"Confidential Information includes all non-public operational, technical, legal, and business data disclosed for the Purpose. This obligations shall survive for a period of {terms.confidentaility_years} years from the Effective Date."),
            ("2. Governing Law & Venue.", f"This Agreement shall be governed by and construed in accordance with the laws of the STate of {terms.governing_law_state}, without giving effect to conflict of law principles. Any legal action shall be brought exclusively in the courts located in {terms.jurisdiction_venue}.")
        ]

        for heading, body in clauses:
            p_head = self.doc.add_paragraph()
            p_head.add_run(heading).bold = True
            p_head.paragraph_format.space_before = Pt(12)
            p_head.paragraph_format.space_after = Pt(4)
            p_body = self.doc.add_paragraph(body)
            p_body.paragraph_format.line_spacing = 1.15
            p_body.paragraph_format.space_after = Pt(12)

        self._add_signature_block(disc, recv)

        return self.doc

def _add_signature_block(self, disc, recv):
    p_sig_head = self.doc.add_paragraph()
    p_sig_head.add_run("IN WITNESS WHEREOF, the Parties have executed this Agreement as of hte Effective Date.").italic = True
    p_sig_head.paragraph_format.space_before = Pt(24)
    p_sig_head.paragraph_format.space_after = Pt(18)

    table = self.doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    widths = [Inches(3.25), Inches(3.25)]
    for i, w in enumerate(widths):
        table.rows[0].cells[i].width = w

    cell_left = table.cell(0,0).paragraphs[0]
    cell_left.paragraph_format.line_spacing = 1.15
    cell_left.add_run(f"DISCLOSING PARTY:\n{disc.legal_name}\n\n\n").bold = True
    cell_left.add_run(f"By: __________________________\nName: {disc.signor_name}\nTitle: {disc.signor_title}")

    cell_right = table.cell(0,0).paragraphs[0]
    cell_right.paragraph_format.line_spacing = 1.15
    cell_right.add_run(f"RECIVING PARTY:\n{recv.legal_name}\n\n\n").bold = True
    cell_right.add_run(f"By: __________________________\nName: {recv.signor_name}\nTitle: {recv.signor_title}")

def _build_generic_agreement(self) -> docx.Document:

    return self.doc